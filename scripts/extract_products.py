from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("/Users/mustafa/Downloads/Труба_гофрированная_характеристики_.xlsx")
SNAPSHOT = ROOT / "data" / "Труба_гофрированная_характеристики_.xlsx"
OUTPUTS = [
    ROOT / "web" / "src" / "data" / "products.json",
    ROOT / "cms" / "src" / "data" / "products.json",
]


def clean(value: object) -> str:
    return str(value or "").strip().replace("\r\n", "\n")


def flag(value: object) -> bool:
    return clean(value).lower() in {"есть", "да", "yes", "true"}


def number(value: object) -> float | None:
    if value in (None, ""):
        return None
    match = re.search(r"-?\d+(?:[.,]\d+)?", clean(value))
    return float(match.group(0).replace(",", ".")) if match else None


def product_image(sku: str) -> str:
    image = f"/images/products/{sku}.webp"
    thumbnail = f"/images/products/{sku}-thumb.webp"
    for asset in (image, thumbnail):
        if not (ROOT / "web" / "public" / asset.lstrip("/")).is_file():
            raise ValueError(f"Generate a unique product image before importing SKU {sku}: {asset}")
    return image


def main() -> None:
    workbook = load_workbook(SOURCE, data_only=True)
    sheet = workbook["Продукция"]
    rows = list(sheet.iter_rows(values_only=True))
    products: list[dict[str, object]] = []

    for row in rows[1:]:
        if not clean(row[0]) or not clean(row[1]):
            continue
        name = clean(row[0])
        sku = re.sub(r"\D", "", clean(row[1]))
        if not sku:
            continue
        material = "ПВХ" if "ПВХ" in name else "ПНД"
        load_class = "Тяжелая" if "Тяжел" in name else "Легкая"
        length_match = re.search(r"длина\s+(\d+)\s*м", name, re.IGNORECASE)
        coil_length = int(length_match.group(1)) if length_match else 0
        color = clean(row[11]).capitalize()
        price = number(row[21]) or 0

        products.append(
            {
                "name": name,
                "slug": f"gofrotruba-{sku}",
                "sku": sku,
                "material": material,
                "loadClass": load_class,
                "description": clean(row[2]),
                "packageType": clean(row[3]),
                "temperature": clean(row[4]),
                "combustibility": clean(row[5]),
                "halogenFree": flag(row[6]),
                "uvResistant": flag(row[7]),
                "frostResistant": bool(clean(row[8])),
                "concretePour": flag(row[9]),
                "protection": clean(row[10]),
                "color": color,
                "innerDiameter": number(row[12]),
                "outerDiameter": number(row[13]),
                "bendRadius": number(row[14]),
                "climate": clean(row[15]),
                "coilWeight": clean(row[16]),
                "coilLength": coil_length,
                "packageLength": number(row[17]),
                "packageWidth": number(row[18]),
                "packageHeight": number(row[19]),
                "compression": clean(row[20]).replace("Н", " Н"),
                "price": price,
                "image": product_image(sku),
                "featured": sku in {"2021022", "2021030", "2021010", "2021018"},
            }
        )

    SNAPSHOT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, SNAPSHOT)
    for output in OUTPUTS:
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(products, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(f"Extracted {len(products)} products")
    print(f"SKUs: {', '.join(str(item['sku']) for item in products)}")


if __name__ == "__main__":
    main()
